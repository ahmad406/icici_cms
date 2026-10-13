import re
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal

# confirm with ICICI
PAYMENT_RECORD_ID = "P"
ANNEXURE_RECORD_ID = "A"

PAYMENT_CODES = ("NEFT", "RTGS", "IMPS", "IFT")
MONTHS = ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC")

PAYMENT_FIELDS = [
	("record_id", "Record_Identifier", 1, True, False),
	("indicator", "Payment Indicator", 30, True, False),
	("ref_no", "Unique Cust Ref No", 30, False, False),
	("party", "Vendor / Beneficiary Code", 500, False, False),
	("party_name", "Name of Beneficiary", 150, True, True),
	("amount", "Instrument Amount", 30, True, False),
	("payment_date", "Payment Date", 20, True, False),
	("cheque_no", "Cheque Number", 6, False, False),
	("debit_account", "Debit Account No", 30, True, False),
	("bene_account", "Beneficiary Bank A/c", 30, True, False),
	("ifsc", "IFSC CODE", 500, True, False),
	("bank_name", "Beneficiary Bank Name", 500, True, True),
	("addr1", "Beneficiary Mailing Address1", 100, False, True),
	("addr2", "Beneficiary Mailing Address2", 100, False, True),
	("addr3", "Beneficiary Mailing Address3", 100, False, True),
	("city", "Beneficiary city", 100, False, True),
	("zip", "Beneficiary Zip", 10, False, False),
	("narration", "Debit narration", 100, False, True),
	("print_location", "Print Location", 150, False, True),
	("payable_location", "Payable Location", 150, False, True),
	("fiscal_year", "Fiscal Year", 500, False, False),
	("company", "Company Code", 500, False, True),
	("email", "Email id", 500, False, False),
	("mobile", "Beneficiary Mobile No", 500, False, False),
	("aadhar", "AADHAR Number", 500, False, False),
]

ANNEXURE_FIELDS = [
	("record_id", "ANNEX_REC_ID", 1, False, False),
	("invoice_no", "Invoice No", 500, False, True),
	("invoice_date", "invoice Date", 500, False, False),
	("gross", "Gross Amount", 500, False, False),
	("tds", "TDS", 500, False, False),
	("deductions", "Deductions", 500, False, False),
	("net", "Net Amount", 500, False, False),
]


def clean(value):
	if value is None:
		return ""
	return re.sub(r"\s+", " ", str(value).replace("|", " ")).strip()


def fmt_amount(value):
	return str(Decimal(str(value or 0)).quantize(Decimal("0.01"), ROUND_HALF_UP))


def fmt_date(value):
	if not value:
		return ""
	if isinstance(value, str):
		value = datetime.strptime(value[:10], "%Y-%m-%d")
	if isinstance(value, (date, datetime)):
		return f"{value.day:02d}-{MONTHS[value.month - 1]}-{value.year}"
	return ""


def get_payment_code(mode_of_payment):
	for token in re.findall(r"[A-Z]+", (mode_of_payment or "").upper()):
		if token in PAYMENT_CODES:
			return token
	return ""


def _build(fields, data):
	values, errors = [], []
	for key, label, max_len, mandatory, trim in fields:
		val = clean(data.get(key))
		if mandatory and not val:
			errors.append(f"{label} is missing")
		if len(val) > max_len:
			if trim:
				val = val[:max_len]
			else:
				errors.append(f"{label} is longer than {max_len} characters ({val})")
		values.append(val)
	return "|".join(values), errors


def build_payment_line(data):
	data = {**data, "record_id": PAYMENT_RECORD_ID}
	return _build(PAYMENT_FIELDS, data)


def build_annexure_line(data):
	data = {**data, "record_id": ANNEXURE_RECORD_ID}
	return _build(ANNEXURE_FIELDS, data)


def build_excel(rows):
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Font

	wb = Workbook()
	pay = wb.active
	pay.title = "PUSH_Payment"
	pay.append([f[1] for f in PAYMENT_FIELDS])
	ann = wb.create_sheet("PUSH_Annexure")
	ann.append(["Unique Cust Ref No"] + [f[1] for f in ANNEXURE_FIELDS])
	for ref_no, pay_values, ann_values in rows:
		pay.append(pay_values)
		for v in ann_values:
			ann.append([ref_no] + v)
	for ws in (pay, ann):
		for cell in ws[1]:
			cell.font = Font(bold=True)
		ws.freeze_panes = "A2"
	buf = BytesIO()
	wb.save(buf)
	return buf.getvalue()
